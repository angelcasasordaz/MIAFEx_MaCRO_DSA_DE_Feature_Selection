"""Validated cache-only reporting and append-only report version publication."""
import argparse
from contextlib import contextmanager
from dataclasses import dataclass
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import sys
import tempfile
import time

import numpy as np


def framework():
    # main_best.py may be running as __main__; reuse that exact module.
    main = sys.modules.get("__main__")
    if getattr(main, "__file__", "").endswith("main_best.py"):
        return main
    import main_best
    return main_best


def sha256(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


@contextmanager
def report_stage(label):
    """Make long cache-only operations visible even in a buffered IDE console."""
    started = time.perf_counter()
    print(f'[report] {label} ...', flush=True)
    try:
        yield
    except BaseException:
        print(f'[report] {label} failed after {time.perf_counter() - started:.1f}s', flush=True)
        raise
    else:
        print(f'[report] {label} complete ({time.perf_counter() - started:.1f}s)', flush=True)


@contextmanager
def report_guard(destinations):
    """Reject scientific execution and filesystem mutations outside replica folders."""
    state = {"active": True, "optimization_calls": 0}
    forbidden = {"_run_single", "execute_pending_runs", "build_optimizer",
                 "configure_compute_backend", "resolve_execution_config", "save_cache", "save_combination",
                 "train_miafex", "extract_miafex_features", "resolve_miafex_csv",
                 "read_miafex_csv", "load_miafex_csv", "load_miafex_feature_data", "get_dataset",
                 "run_single", "run_single_parallel_task", "start_gpu_request_service",
                 "load_dataset"}
    blocked_codes = {}
    def profile(frame, event, arg):
        if event != 'call':
            return
        # Code metadata is immutable. Inspect it once instead of re-triggering
        # object.__getattr__ audit events on every library function call.
        code = frame.f_code
        blocked = blocked_codes.get(code)
        if blocked is None:
            name = code.co_name
            blocked = (name in forbidden or (name == 'solve' and 'mealpy' in code.co_filename)
                       or (name in {'fit', 'partial_fit'} and any(lib in code.co_filename for lib in ('mafese', 'sklearn', 'torch'))))
            blocked_codes[code] = blocked
        if blocked:
            state["optimization_calls"] += 1
            raise RuntimeError(f"REPORT-ONLY blocked scientific execution: {code.co_name}")
    def allowed(path):
        if isinstance(path, int):
            return False
        resolved = Path(os.fsdecode(path)).resolve()
        return any(resolved == base or base in resolved.parents for base in destinations) and resolved.suffix.lower() not in {'.pkl', '.pickle', '.ckpt', '.pth', '.pt', '.npy', '.npz', '.pdf'}
    mutation_events = frozenset({'open', 'os.remove', 'os.rmdir', 'os.mkdir', 'os.chmod',
                                'os.utime', 'os.truncate', 'shutil.rmtree', 'os.rename',
                                'os.link', 'os.symlink', 'subprocess.Popen', 'os.system', 'os.fork'})
    def audit(event, args):
        if not state["active"] or event not in mutation_events:
            return
        targets = []
        def at(path, dir_fd):
            # shutil.rmtree uses descriptor-relative unlink/rmdir on Linux.
            if dir_fd is not None and dir_fd != -1 and not os.path.isabs(path):
                return Path(os.readlink(f"/proc/self/fd/{dir_fd}")) / os.fsdecode(path)
            return path
        if event == "open":
            path, mode, flags = args
            if (flags & (os.O_WRONLY | os.O_RDWR | os.O_CREAT | os.O_TRUNC | os.O_APPEND)):
                targets = [path]
        elif event in {"os.remove", "os.rmdir"}:
            targets = [at(args[0], args[1])]
        elif event in {"os.mkdir", "os.chmod"}:
            targets = [at(args[0], args[2])]
        elif event in {"os.utime", "os.truncate", "shutil.rmtree"}:
            targets = [args[0]]
        elif event in {"os.rename", "os.link", "os.symlink"}:
            targets = ([at(args[0], args[2]), at(args[1], args[3])]
                       if event in {'os.rename', 'os.link'} else [args[0], at(args[1], args[2])])
        elif event in {"subprocess.Popen", "os.system", "os.fork"}:
            raise RuntimeError(f"REPORT-ONLY blocked process execution: {event}")
        if any(not allowed(path) for path in targets):
            raise PermissionError(f"REPORT-ONLY blocked mutation outside replica output: {event} {targets}")
    previous = sys.getprofile()
    previous_bytecode = sys.dont_write_bytecode
    sys.dont_write_bytecode = True
    sys.addaudithook(audit)
    sys.setprofile(profile)
    try:
        yield state
    finally:
        state["active"] = False
        sys.setprofile(previous)
        sys.dont_write_bytecode = previous_bytecode


@dataclass
class CompletedReport:
    args: argparse.Namespace
    results: dict
    indexed: dict
    datasets: list
    classifiers: list
    algorithms: list
    metrics: list
    signature: str
    sources: dict
    identities: dict = None

    @property
    def exp_tag(self):
        return f"EXP{self.args.exp_id:03d}"


def safe_path(path):
    """Reject redirects, traversal and non-directory ancestors before any write."""
    path = Path(path).absolute()
    if '..' in path.parts or any(p.is_symlink() for p in (path, *path.parents)):
        raise ValueError(f"Unsafe reporting path: {path}")
    if any(p.exists() and not p.is_dir() for p in path.parents):
        raise ValueError(f"Non-directory reporting ancestor: {path}")
    return path


def _read_source(path, root, sources, loaded):
    path = safe_path(path)
    if not path.is_file():
        raise FileNotFoundError(f"Missing completed cache: {path}; no recomputation allowed")
    if path not in loaded:
        before = sha256(path)
        if path.suffix == '.json':
            value = json.loads(path.read_text(encoding='utf-8'))
        else:
            value = framework().load_cache(str(path))
        if sha256(path) != before:
            raise ValueError(f"Cache content changed during read: {path}")
        sources[str(path.relative_to(root))] = before
        loaded[path] = value
    return loaded[path]


def _matches_request(identity, args):
    """Select stored science for this report, without rebuilding an execution identity.

    Historical implementation/package/feature fingerprints remain authoritative.
    Reporting never claims compatibility with the current optimizer code, and
    never loads datasets, builds optimizers, imports runs or repairs caches.
    """
    if not isinstance(identity, dict) or identity.get('schema') != 2:
        raise ValueError('Invalid scientific cache identity')
    required = {'dataset', 'dataset_source', 'classifier', 'optimizer', 'transfer_function',
                'runs', 'epochs', 'pop_size', 'objective', 'prepared_features', 'partition',
                'seeds', 'optimizer_parameters', 'optimizer_implementation',
                'binary_representation', 'wrapper_revision', 'wrapper_sources', 'packages'}
    if not required <= identity.keys():
        raise ValueError('Incomplete scientific cache identity')
    if any(identity[key] != getattr(args, key) for key in ('dataset_source', 'runs', 'epochs', 'pop_size')):
        return False
    partition, seeds, objective = identity['partition'], identity['seeds'], identity['objective']
    mode = 'prepared_train_test_v1' if args.dataset_source == 'miafex' else 'mafese_split_v1'
    if (partition.get('mode') != mode or partition.get('test_size') != args.test_size
            or partition.get('random_state') != args.random_state
            or seeds.get('base') != args.seed_base
            or seeds.get('run_seeds') != [args.seed_base + r for r in range(args.runs)]
            or objective.get('name') != 'AS'
            or objective.get('mode') != 'minimize_metric_loss_plus_feature_ratio_v1'
            or objective.get('weights') != [.9, .1]):
        return False
    method = identity['optimizer']
    parameters = {}
    if method in {'DSADE', 'MaCRO-DE'}:
        parameters = {key: getattr(args, key) for key in
                      ('dsade_beta_min', 'dsade_beta_max', 'dsade_pcr', 'dsade_mahal_q')}
    elif method == 'MaCRO-DE-t':
        parameters = {'wf': .5, 'cr': .9, 'dsade_mahal_q': args.dsade_mahal_q}
    elif method == 'MaCRO-DE-t-v2':
        parameters = {
            'beta_min': args.macro_de_t_v2_beta_min,
            'beta_max': args.macro_de_t_v2_beta_max,
            'mahalanobis_q': args.macro_de_t_v2_mahal_q,
            'pcr_policy': '0.1 + 0.25 * (1.0 - dM)',
        }
    return (method in args.optimizers and identity['classifier'] in args.estimators
            and identity['transfer_function'] in args.transfer_functions
            and identity['optimizer_parameters'] == parameters)


def _validate_row(row, identity, args, context):
    from reporting.paper_tables import METRICS
    m = framework()
    if not isinstance(row, dict) or str(row.get('Estimator', '')).lower() != identity['classifier']:
        raise ValueError(f'Classifier identity mismatch: {context}')
    if row.get('CompletedRuns') != args.runs:
        raise ValueError(f'Incomplete cache: {context}; completed run count mismatch')
    present = [metric for metric in METRICS if metric.run_key in row]
    if not present:
        raise ValueError(f'No available run metrics: {context}')
    # Stored order is the seed/run position used for cross-dataset manuscript means.
    ids = list(row.get('CompletedRunIDs', range(args.runs)))
    if ids != list(range(args.runs)):
        raise ValueError(f'Incomplete or unordered run IDs: {context}')
    for metric in present:
        values = np.asarray(row[metric.run_key], dtype=float)
        if values.shape != (args.runs,) or not np.isfinite(values).all():
            raise ValueError(f'Invalid run values: {context}/{metric.run_key}')
        mean_key = metric.run_key.replace('Runs', 'Mean')
        if not np.isclose(row.get(mean_key, np.nan), values.mean(), rtol=1e-12, atol=1e-12):
            raise ValueError(f'Cached summary disagrees with runs: {context}/{mean_key}')
    curves = row.get('CurvesAll')
    if curves is not None:
        if len(curves) != args.runs:
            raise ValueError(f'Incomplete convergence run array: {context}')
        for i, curve in enumerate(curves):
            final = row['FitRuns'][i] if 'FitRuns' in row else None
            m.validate_convergence_curve(curve, args.epochs, final)
        evidence = row.get('ConvergenceRuns', {})
        if not isinstance(evidence, dict) or set(evidence) - set(ids):
            raise ValueError(f'Invalid convergence evidence: {context}')
        for run, metadata in evidence.items():
            m.validate_convergence_metadata(metadata, curves[run], args.epochs)
    mean = np.asarray(row.get('Curve', []), dtype=float)
    if mean.size:
        m.validate_convergence_curve(mean, args.epochs)
        if curves is None or not np.allclose(mean, m.pad_mean_curves(curves, args.epochs), rtol=1e-12, atol=1e-15):
            raise ValueError(f'Invalid cached mean convergence: {context}')


def load_completed_cache(args):
    """Read final v2 envelopes or SHA-256-bound migrated final references only.

    The selected EXP is the sole source. Progress caches, cross-EXP fallback,
    migration, convergence repair and all scientific execution are excluded.
    Ambiguous scientific identities and partial requested grids fail closed.
    """
    from reporting.paper_tables import METRICS
    import scientific_cache
    m = framework()
    args = argparse.Namespace(**vars(args))
    if not isinstance(args.exp_id, int) or args.exp_id < 0 or args.runs < 1 or args.epochs < 1:
        raise ValueError('A nonnegative EXP ID and positive run/epoch counts are required')
    m.validate_selection_options(args)
    args.optimizers = m.resolve_optimizers(args)
    args.estimators = list(dict.fromkeys(args.estimators))
    args.transfer_functions = list(dict.fromkeys(args.transfer_functions))
    if not args.optimizers or not args.estimators or not args.transfer_functions:
        raise ValueError('Select optimizers, classifiers and transfer functions')
    root = safe_path(args.output_root)
    cache = safe_path(root / 'Results' / f'EXP{args.exp_id:03d}' / 'cache')
    inventory = safe_path(cache / 'combinations_v2')
    if not inventory.is_dir():
        raise FileNotFoundError(f'Missing completed cache: {inventory}; no recomputation allowed')
    sources, loaded, candidates = {}, {}, {}
    for path in sorted(inventory.glob('*_results.*')):
        if path.suffix not in {'.pkl', '.json'}:
            continue
        entry = _read_source(path, root, sources, loaded)
        identity = entry.get('identity') if isinstance(entry, dict) else None
        if not _matches_request(identity, args):
            continue
        digest = scientific_cache.identity_digest(identity)
        if path.name != f'{digest}_results{path.suffix}':
            raise ValueError(f'Scientific identity/filename mismatch: {path}')
        key = tuple(identity[field] for field in ('dataset', 'classifier', 'optimizer', 'transfer_function'))
        candidates.setdefault(key, []).append((path, entry, digest))
    if args.dataset_source == 'mafese':
        datasets = m.resolve_mafese_dataset_names(args)
    else:
        datasets = [args.dataset_name] if args.dataset_name else args.miafex_datasets
        if datasets is None:
            # Directory metadata only: no image or feature contents are opened.
            features = Path(args.feature_dataset_root)
            names = {key[0] for key in candidates}
            if features.is_dir():
                names.update(p.name for p in features.iterdir() if p.is_dir()
                             and all((p / f'{split}_features.csv').is_file() for split in ('train', 'test')))
            datasets = sorted(names, key=str.lower)
    datasets = list(dict.fromkeys(datasets))
    if not datasets:
        raise FileNotFoundError('No completed dataset caches; no recomputation allowed')
    if any(not isinstance(ds, str) or not ds or '/' in ds or '\\' in ds or ds in {'.', '..'} for ds in datasets):
        raise ValueError('Unsafe dataset identity')
    results = {dataset: {} for dataset in datasets}
    indexed, identities, algorithms, shared, features = {}, {}, [], None, {}
    for dataset in datasets:
        for classifier in args.estimators:
            for method in args.optimizers:
                for transfer in args.transfer_functions:
                    key = dataset, classifier, method, transfer
                    options = candidates.get(key, [])
                    if not options:
                        raise FileNotFoundError(f'Missing completed cache: {key}; no recomputation allowed')
                    if len({digest for _, _, digest in options}) != 1:
                        raise ValueError(f'Ambiguous scientific identity: {key}')
                    # A final envelope takes precedence over a redundant migration reference.
                    path, entry, digest = min(options, key=lambda option: option[0].suffix != '.pkl')
                    identity = entry['identity']
                    if path.suffix == '.pkl':
                        row = entry.get('row')
                    else:
                        origins = [origin for origin in entry.get('sources', [])
                                   if str(origin.get('file', '')).endswith('_results.pkl')]
                        if len(origins) != 1:
                            raise ValueError(f'Missing/ambiguous migrated final source: {path}')
                        origin = origins[0]
                        name = origin['file']
                        if Path(name).name != name or '/' in name or '\\' in name:
                            raise ValueError(f'Unsafe migrated source filename: {name}')
                        original = safe_path(cache / name)
                        payload = _read_source(original, root, sources, loaded)
                        if sources[str(original.relative_to(root))] != origin['sha256']:
                            raise ValueError(f'Original payload changed: {name}')
                        row = payload.get(origin['label']) if isinstance(payload, dict) else None
                        parsed = m.parse_result_label(origin['label'], args)
                        if (parsed['method'] != method or parsed['estimator'] not in {'', classifier}
                                or parsed['transfer_function'] not in {'', transfer}):
                            raise ValueError(f'Migrated label identity mismatch: {origin["label"]}')
                    _validate_row(row, identity, args, str(path))
                    common = {field: identity[field] for field in
                              ('schema', 'dataset_source', 'objective', 'partition', 'seeds',
                               'wrapper_revision', 'wrapper_sources', 'packages')}
                    if shared is not None and common != shared:
                        raise ValueError('Mixed scientific configurations within reporting grid')
                    shared = common
                    if dataset in features and features[dataset] != identity['prepared_features']:
                        raise ValueError(f'Mixed prepared feature identities: {dataset}')
                    features[dataset] = identity['prepared_features']
                    algorithm = m.optimizer_acronym(method) + (f'_{transfer.upper()}' if len(args.transfer_functions) > 1 else '')
                    label = m.build_alg_label(method, transfer, classifier, len(args.transfer_functions) > 1, True)
                    indexed[dataset, classifier, algorithm] = row
                    results[dataset][label] = row
                    identities[digest] = identity
                    if algorithm not in algorithms:
                        algorithms.append(algorithm)
    metrics = [metric for metric in METRICS if any(metric.run_key in row for row in indexed.values())]
    for metric in metrics:
        if any(metric.run_key not in row for row in indexed.values()):
            raise ValueError(f'Incomplete metric grid: {metric.run_key}')
    # Only selected files are provenance; inventory metadata was read but never used as observations.
    selected_paths = {str(path.relative_to(root)) for key in candidates if key[0] in datasets
                      for path, _, _ in candidates[key]}
    selected_paths.update(path for path in sources if Path(path).parent == cache.relative_to(root))
    sources = {path: digest for path, digest in sources.items() if path in selected_paths}
    signature = scientific_cache.identity_digest(sorted(identities))
    return CompletedReport(args, results, indexed, datasets, args.estimators, algorithms,
                           metrics, signature, sources, identities)


def report_versions(root, exp_id):
    """Inspect both trees. Unpaired, redirected or malformed versions are errors."""
    if not isinstance(exp_id, int) or exp_id < 0:
        raise ValueError('A nonnegative user-selected EXP ID is required')
    root = safe_path(root)
    versions = []
    for kind in ('Figures', 'Results'):
        parent = safe_path(root / kind / f'EXP{exp_id:03d}')
        found = set()
        if parent.exists():
            for path in parent.iterdir():
                match = re.fullmatch(r'full_rep([1-9][0-9]*)', path.name)
                if not match:
                    if re.fullmatch(r'full_rep[0-9]+', path.name):
                        raise ValueError(f"Noncanonical report version: {path}")
                    continue
                safe_path(path)
                if not path.is_dir():
                    raise ValueError(f"Conflicting report version: {path}")
                for child in path.rglob('*'):
                    if child.is_symlink():
                        raise ValueError(f"Symlink within report: {child}")
                manifest_path = path / 'validation.json'
                if manifest_path.is_file():
                    manifest = json.loads(manifest_path.read_text(encoding='utf-8'))
                    if (manifest.get('experiment_id', exp_id) != exp_id
                            or manifest.get('report_version', int(match[1])) != int(match[1])):
                        raise ValueError(f'Conflicting report manifest identity: {manifest_path}')
                found.add(int(match[1]))
        versions.append(found)
    if versions[0] != versions[1]:
        raise ValueError(f"Incomplete Figures/Results report pairs: {versions}; existing files preserved")
    return sorted(versions[0])


def next_report_version(root, exp_id):
    return max(report_versions(root, exp_id), default=0) + 1


def _rename_new(source, destination):
    """Atomic, no-replace directory publication (Linux); fail closed elsewhere."""
    import ctypes
    libc = ctypes.CDLL(None, use_errno=True)
    rename = getattr(libc, 'renameat2', None)
    if rename is None:
        raise RuntimeError('Atomic no-replace publication requires renameat2')
    if rename(-100, os.fsencode(source), -100, os.fsencode(destination), 1):
        error = ctypes.get_errno()
        raise OSError(error, os.strerror(error), str(destination))


@contextmanager
def staged_version(root, exp_id):
    """Serialize allocators, stage both trees, then publish a matching pair.

    A crash between the two renames leaves an unpaired version: the next attempt
    fails safely for manual inspection. Temporary names never consume a number.
    """
    root = safe_path(root)
    next_report_version(root, exp_id)  # Preflight before directory creation.
    parents = [safe_path(root / k / f'EXP{exp_id:03d}') for k in ('Figures', 'Results')]
    for parent in parents:
        parent.mkdir(parents=True, exist_ok=True)
    lock = parents[1] / '.report-allocation.lock'
    try:
        lock.mkdir()
    except FileExistsError as exc:
        raise ValueError(f'Report allocation is locked; inspect before retrying: {lock}') from exc
    stages, published = [], []
    try:
        version = next_report_version(root, exp_id)
        finals = [p / f'full_rep{version}' for p in parents]
        for parent in parents:
            stages.append(Path(tempfile.mkdtemp(prefix='.report-staging-', dir=parent)))
        yield version, stages[0], stages[1], finals
        for source, target in zip(stages, finals):
            safe_path(target)
            inode = source.stat().st_ino
            _rename_new(source, target)
            published.append((source, target, inode))
    except BaseException:
        for source, target, inode in reversed(published):
            if not target.is_symlink() and target.stat().st_ino == inode:
                _rename_new(target, source)
        raise
    finally:
        for stage in stages:
            if stage.exists() and not stage.is_symlink():
                shutil.rmtree(stage)
        lock.rmdir()


def _validate_artifacts(figures, results, required):
    from openpyxl import load_workbook
    from PIL import Image
    for relative in required:
        if not (results / relative).is_file():
            raise ValueError(f'Missing required report output: {relative}')
    hashes = {}
    for base in (figures, results):
        for path in base.rglob('*'):
            if path.is_symlink():
                raise ValueError(f'Redirected generated output: {path}')
            if not path.is_file():
                continue
            if path.suffix.lower() in {'.pdf', '.pkl', '.pickle', '.ckpt', '.pth', '.pt', '.npy', '.npz'} or path.stat().st_size == 0:
                raise ValueError(f'Invalid generated output: {path}')
            if path.suffix == '.xlsx':
                wb = load_workbook(path, data_only=False)
                try:
                    if not wb.sheetnames or any(ws.max_row < 1 for ws in wb):
                        raise ValueError(f'Empty workbook: {path}')
                    if path.name.startswith('Paper_Tables'):
                        from reporting.paper_tables import validate_plain_workbook
                        validate_plain_workbook(wb)
                finally:
                    wb.close()
            if path.suffix == '.png':
                with Image.open(path) as png:
                    if png.format != 'PNG' or any(abs(v - 600) > .1 for v in png.info.get('dpi', (0, 0))):
                        raise ValueError(f'Expected 600 dpi PNG: {path}')
                    png.verify()
            hashes[('Figures/' if base == figures else 'Results/') + str(path.relative_to(base))] = sha256(path)
    return hashes


def generate_outputs(report, figures, results):
    """Reuse framework Excel exporters without calling its experiment dispatcher."""
    from reporting import figures as plotting, statistics, paper_tables
    m, args = framework(), report.args
    figures.mkdir(parents=True, exist_ok=True)
    results.mkdir(parents=True, exist_ok=True)
    tag = report.exp_tag
    required = [f'Global_Results_{tag}.xlsx', f'Statistical_Results_{tag}.xlsx', f'Paper_Tables_{tag}.xlsx',
                'statistics/statistical_summary.csv', 'statistics/pairwise_wilcoxon_holm.csv',
                'statistics/statistical_report.txt']
    with report_stage(f'{tag}: Global Results Excel'):
        m.export_global_excel(report.results, report.datasets, str(results / required[0]))
    with report_stage(f'{tag}: Statistical Results Excel'):
        m.export_statistical_excel(report.results, report.datasets, args.optimizers, args, str(results / required[1]))
    with report_stage(f'{tag}: paper tables (format, serialize, validate)'):
        paper_tables.export_indexed_tables(report.indexed, report.datasets, report.algorithms, report.classifiers,
                                          report.metrics, results / required[2], title=f'{tag} MIAFEx')
    skipped = [{'output': f'{metric.name} tables/figures', 'reason': f'{metric.run_key} unavailable in every cache row'}
               for metric in paper_tables.METRICS if metric not in report.metrics]
    if any(metric.run_key == 'FitRuns' for metric in report.metrics):
        name = f'Full_Friedman_Analysis_{tag}.xlsx'
        with report_stage(f'{tag}: Friedman Excel'):
            m.export_friedman_analysis(report.results, report.datasets, args.optimizers, args, str(results / name))
        required.append(name)
    else:
        skipped.append({'output': 'Framework fitness Friedman Excel', 'reason': 'Requires available FitRuns'})
    with report_stage(f'{tag}: publication figures at 600 dpi'):
        skipped.extend(plotting.generate(report, figures))
    with report_stage(f'{tag}: statistical analysis and figures'):
        skipped.extend(statistics.export(report, figures / 'statistics', results / 'statistics'))
    return required, skipped


def run_full_figures(args):
    """Regenerate the FULL figure tree from validated caches; never export tables."""
    from reporting import figures, statistics
    m = framework()
    with report_stage('Validate completed caches for FULL figures'), report_guard(()) as preflight:
        report = load_completed_cache(args)
    root = safe_path(args.output_root)
    figure_root = safe_path(root / 'Figures' / report.exp_tag)
    destination_root = safe_path(getattr(args, 'report_output_root', None) or root)
    destination = safe_path(destination_root / 'Figures' / report.exp_tag / 'full')
    destination.mkdir(parents=True, exist_ok=True)
    # Retain filenames while relocating the old flat figure tree exactly once.
    moved = []
    for path in figure_root.glob('*'):
        if not path.is_file() or path.suffix.lower() not in {'.png', '.pdf'}:
            continue
        subdirectory = ''
        if any(path.stem == f'{family}_{dataset}_{classifier}'
               for family in ('01_resultados_clasificador', '02_radar', '03_features_runtime', '05_convergence')
               for dataset in report.datasets for classifier in report.classifiers):
            subdirectory = 'individual'
        elif path.stem.startswith(('generic_average_rank', 'generic_block_distribution',
                                   'generic_holm_heatmap', 'generic_reference_comparisons')):
            subdirectory = 'statistics'
        elif not path.stem.startswith(tuple(f'{n:02d}_' for n in range(1, 10))):
            continue
        # An alternate destination must not move historical source figures.
        if destination_root != root:
            continue
        target = safe_path(destination / subdirectory / path.name)
        if target.exists():
            raise ValueError(f'Figure relocation would replace an existing file: {target}')
        target.parent.mkdir(parents=True, exist_ok=True)
        path.rename(target)
        moved.append(str(target.relative_to(destination)))
    with report_stage('Generate representative FULL figures'), report_guard((destination,)) as guard:
        generated = m.generate_seven_global_charts(
            m.generate_plot_dataframe(report.results, report.args), report.results,
            str(destination), report.args.optimizers, report.args,
        )
        metric = next((metric for metric in report.metrics if metric.run_key == 'F1Runs'), report.metrics[0])
        _, matrix = statistics.matched_block_matrix(
            report.indexed, report.datasets, report.classifiers, report.algorithms, metric,
        )
        analysis = statistics.analyze(matrix, report.algorithms, higher_is_better=metric.best_mode == 'max')
        stats_dir = destination / 'statistics'
        stats_dir.mkdir(exist_ok=True)
        for stem, fig in figures.statistical_figures(analysis, report.algorithms, metric.name):
            figures.save_png(fig, stats_dir / f'{stem}.png')
            generated.append(f'statistics/{stem}.png')
        if any(sha256(root / path) != digest for path, digest in report.sources.items()):
            raise ValueError('Source cache changed during figure generation')
    print(f'[report] FULL figures: {destination}; aggregation={m.PLOT_RUN_AGGREGATION}; '
          f'optimization_calls={guard["optimization_calls"] + preflight["optimization_calls"]}', flush=True)
    return {'figures': str(destination), 'generated': generated, 'moved': moved,
            'optimization_calls': guard['optimization_calls'] + preflight['optimization_calls']}


def run_report(args):
    """Append one validated full_repN for the explicitly selected MIAFEx EXP."""
    started = time.perf_counter()
    print(f'[report] Starting cache-only EXP{args.exp_id:03d}', flush=True)
    # Import dependencies before the strict mutation guard is active.
    from reporting import figures, statistics, paper_tables
    root = safe_path(args.output_root)
    with report_stage('Validate original completed caches'), report_guard(()) as preflight:
        report = load_completed_cache(args)
    with report_stage('Hash protected historical outputs'):
        previous = {str(p.relative_to(root)): (sha256(p), p.stat().st_mtime_ns)
                    for kind in ('Figures', 'Results')
                    for p in (root / kind / report.exp_tag).rglob('*') if p.is_file()}
    destination_root = safe_path(getattr(args, 'report_output_root', None) or root)
    print(f'[report] Source: {root}; destination: {destination_root}', flush=True)
    with staged_version(destination_root, args.exp_id) as (version, fig, res, finals):
        old_tempdir = tempfile.tempdir
        tempfile.tempdir = str(res)
        try:
            with report_guard((fig, res)) as guard:
                required, skipped = generate_outputs(report, fig, res)
                entry = {'mode': 'full', 'subdirectory': '.',
                         'datasets': report.datasets, 'classifiers': report.classifiers,
                         'algorithms': report.algorithms, 'metrics': [m.run_key for m in report.metrics],
                         'source_directory': str(root / 'Results' / report.exp_tag / 'cache'),
                         'completed_runs': sorted({r['CompletedRuns'] for r in report.indexed.values()}),
                         'cache_identity': report.signature, 'scientific_identities': report.identities,
                         'source_cache_sha256': report.sources, 'skipped_outputs': skipped}
                with report_stage('Validate every generated artifact'):
                    hashes = _validate_artifacts(fig, res, required)
                with report_stage('Verify protected historical hashes'):
                    if any(not (root / path).is_file() or sha256(root / path) != digest
                           or (root / path).stat().st_mtime_ns != mtime
                           for path, (digest, mtime) in previous.items()):
                        raise ValueError('Protected source/report changed during reporting')
                    if any(sha256(root / path) != digest for path, digest in report.sources.items()):
                        raise ValueError('Source cache changed during reporting')
                manifest = {'experiment_id': args.exp_id, 'report_version': version,
                            'report_version_is_scientific_repetition': False, 'reports': [entry],
                            'figures_destination': str(finals[0]), 'results_destination': str(finals[1]),
                            'outputs_sha256': hashes, 'required_results': required,
                            'optimization_calls': guard['optimization_calls'] + preflight['optimization_calls'],
                            'dpi': 600, 'pdfs_generated': 0, 'protected_files_unchanged': True}
                for base in (fig, res):
                    (base / 'validation.json').write_text(json.dumps(manifest, indent=2), encoding='utf-8')
        finally:
            tempfile.tempdir = old_tempdir
    print(f'[report] Published {report.exp_tag} full_rep{version} in {time.perf_counter() - started:.1f}s: '
          f'{finals[0]} and {finals[1]}', flush=True)
    return manifest
