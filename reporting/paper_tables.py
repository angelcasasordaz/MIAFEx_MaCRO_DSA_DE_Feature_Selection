"""Manuscript Excel tables from MIAFEx results, without experiment execution.

Overall statistics describe equal-weight dataset averages at each run position,
not pooled observations or averages of dataset standard deviations. Matching
datasets must have equal run counts for each algorithm/classifier/metric; counts
may differ between those groups. Accuracy is converted from percent to 0–1;
other metrics retain their stored units. Non-finite statistics follow _run_stats.
"""
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
import tempfile
import math

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.utils import get_column_letter


STATS = ("Best", "Worst", "Mean", "Std")


@dataclass(frozen=True)
class Metric:
    name: str
    run_key: str
    best_mode: str = "max"
    scale: float = 1.0
    unit: str = "0–1"


METRICS = (
    Metric("Accuracy", "AccRuns", scale=100.0),
    Metric("F1-Score", "F1Runs"),
    Metric("Precision", "PSRuns"),
    Metric("Recall", "RSRuns"),
    Metric("Fitness", "FitRuns", "min", unit="fitness"),
    Metric("Features", "FeatRuns", "min", unit="count"),
    Metric("Time", "TimeRuns", "min", unit="seconds"),
)


def calculate_tables(indexed, algorithms, layouts, run_stats):
    """Layouts map sheet names to (header, subheader, datasets, classifier, metric).

Each column averages its datasets by run position before calling the framework's
unchanged statistical reducer. A single dataset therefore uses its original runs.
"""
    tables = {}
    for name, columns in layouts.items():
        rows, means = [], []
        for algorithm in algorithms:
            statistics = []
            for _, _, datasets, classifier, metric in columns:
                runs = []
                for dataset in datasets:
                    key = (dataset, classifier, algorithm)
                    if key not in indexed or metric.run_key not in indexed[key]:
                        raise ValueError(f"Missing manuscript runs: {key}/{metric.run_key}")
                    values = np.asarray(indexed[key][metric.run_key], dtype=float).ravel()
                    if not values.size:
                        raise ValueError(f"Empty manuscript runs: {key}/{metric.run_key}")
                    runs.append(values / metric.scale)
                if len({len(values) for values in runs}) != 1:
                    raise ValueError(f"Unequal dataset run counts: {algorithm}/{classifier}/{metric.name}")
                statistics.append(run_stats(np.mean(np.stack(runs), axis=0), metric.best_mode))
            rows.extend([[item[stat] for item in statistics] for stat in STATS])
            means.append([item["Mean"] for item in statistics])
        tables[name] = (np.asarray(rows), np.asarray(means))
    return tables


def make_workbook(tables, algorithms, layouts, *, title, description,
                  dataset_heading="Dataset", algorithm_labels=None, notes=None):
    """Shared manuscript formatting, with dimensions derived from each table."""
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title, wb.properties.description = title, description
    for sheet_index, (name, columns) in enumerate(layouts.items()):
        ws = wb.create_sheet(name)
        values, references = tables[name]
        if values.shape != (len(algorithms) * 4, len(columns)):
            raise ValueError(f"Invalid manuscript table dimensions: {name}")
        ws.merge_cells("A1:B1")
        ws["A1"] = "Performance metric" if sheet_index == 0 else dataset_heading
        ws["A2"], ws["B2"] = "Algorithm", "Statistic"
        start = 0
        while start < len(columns):
            end = start + 1
            while end < len(columns) and columns[end][0] == columns[start][0]:
                end += 1
            if end - start > 1:
                ws.merge_cells(start_row=1, start_column=start+3, end_row=1, end_column=end+2)
            ws.cell(1, start+3, columns[start][0])
            start = end
        for col, column in enumerate(columns, 3):
            ws.cell(2, col, column[1])
        ws["A1"].comment = Comment((notes or {}).get(name, description), "Paper table reporting")
        for i, opt in enumerate(algorithms):
            first = 3+i*4
            ws.merge_cells(start_row=first, start_column=1, end_row=first+3, end_column=1)
            ws.cell(first, 1, (algorithm_labels or {}).get(opt, opt))
            for stat_index, stat in enumerate(STATS):
                row = first+stat_index
                ws.cell(row, 2, stat)
                for col, value in enumerate(values[i*4+stat_index], 3):
                    ws.cell(row, col, float(value) if np.isfinite(value) else None).number_format = "0.0000"
        ws.freeze_panes = "C3"
        ws.print_title_rows = "1:2"
        apply_plain_presentation(ws)
        actual = [[ws.cell(5+i*4, col).value for col in range(3, len(columns)+3)]
                  for i in range(len(algorithms))]
        if not np.allclose(np.asarray(actual, dtype=float), references, rtol=1e-12, atol=1e-12, equal_nan=True):
            raise ValueError(f"Worksheet Mean validation failed: {name}")
    validate_plain_workbook(wb)
    return wb


def workbook_bytes(wb, tables, *, temp_dir=None):
    """Round-trip every statistic before returning bytes for the final write."""
    buffer = BytesIO()
    old_tempdir = tempfile.tempdir
    try:
        if temp_dir is not None:
            tempfile.tempdir = str(temp_dir)
        wb.save(buffer)
    finally:
        tempfile.tempdir = old_tempdir
    buffer.seek(0)
    check = load_workbook(buffer, data_only=True)
    try:
        validate_plain_workbook(check)
        if check.sheetnames != list(tables):
            raise ValueError("Serialized manuscript sheets changed")
        for name, (values, _) in tables.items():
            ws = check[name]
            if (ws.max_row, ws.max_column) != (values.shape[0]+2, values.shape[1]+2):
                raise ValueError(f"Serialized manuscript dimensions changed: {name}")
            actual = np.asarray([[ws.cell(row+3, col+3).value for col in range(values.shape[1])]
                                 for row in range(values.shape[0])], dtype=float)
            if not np.allclose(actual, values, rtol=1e-12, atol=1e-12, equal_nan=True):
                raise ValueError(f"Serialized manuscript statistics changed: {name}")
    finally:
        check.close()
    return buffer.getvalue()


def export_paper_tables(results_struct, dataset_names, optimizer_order, args, paths):
    """Export supplied MIAFEx results beside existing workbooks; return the path.

Only metrics present in the supplied run records are selected. Missing grid
entries or metric arrays are errors, rather than silently pooled partial data.
Multiple transfer variants of one algorithm/classifier are likewise ambiguous.
"""
    if getattr(args, "report_only", False):
        raise ValueError("Use export_indexed_tables for cache-only reports")
    import main_best as m
    indexed = {}
    for dataset, records in results_struct.items():
        for label, row in records.items():
            parsed = m.parse_result_label(label, args)
            classifier = str(row.get("Estimator") or parsed["estimator"]).lower()
            if not classifier:
                raise ValueError(f"Missing manuscript classifier: {dataset}/{label}")
            key = dataset, classifier, m.optimizer_acronym(parsed["method"])
            if key in indexed:
                raise ValueError(f"Ambiguous manuscript identity: {key}")
            indexed[key] = row
    if not indexed:
        raise ValueError("No run results for manuscript tables")
    def ordered(preferred, present):
        return list(dict.fromkeys([item for item in preferred if item in present] + list(present)))
    datasets = ordered(dataset_names, [key[0] for key in indexed])
    algorithms = ordered([m.optimizer_acronym(opt) for opt in optimizer_order], [key[2] for key in indexed])
    classifiers = ordered([str(c).lower() for c in args.estimators], [key[1] for key in indexed])
    metrics = [metric for metric in METRICS if any(metric.run_key in row for row in indexed.values())]
    if not metrics:
        raise ValueError("No supported manuscript run metrics")
    target = Path(paths.res_dir) / f"Paper_Tables_{paths.exp_tag}.xlsx"
    return export_indexed_tables(indexed, datasets, algorithms, classifiers, metrics, target,
                                 title=f"{paths.exp_tag} MIAFEx")


def export_indexed_tables(indexed, datasets, algorithms, classifiers, metrics, target, *, title):
    """Export validated optimizer/variant identities using the shared reducer."""
    from reporting.core import framework
    m = framework()
    layouts = {"Overall": [(metric.name, cls.upper(), datasets, cls, metric)
                           for metric in metrics for cls in classifiers]}
    for index, cls in enumerate(classifiers, 1):
        layouts[f"Datasets_{index}"] = [(f"{ds} | {cls.upper()}", metric.name, [ds], cls, metric)
                                      for ds in datasets for metric in metrics]
    tables = calculate_tables(indexed, algorithms, layouts, m._run_stats)
    counts = sorted({np.asarray(row[metric.run_key]).size for row in indexed.values() for metric in metrics})
    description = (
        f"{title}; run counts per dataset/algorithm/classifier/metric: {counts}. "
        "Overall: matching run positions averaged equally across datasets before Best/Worst/Mean/Std. "
        "Std: sample ddof=1, zero for one finite run; non-finite observations excluded by _run_stats. "
        "Units: " + "; ".join(f"{metric.name}: {metric.unit}" for metric in metrics)
    )
    wb = make_workbook(tables, algorithms, layouts, title=f"{title} — Paper Tables", description=description)
    payload = workbook_bytes(wb, tables)
    target = Path(target)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(payload)
    return str(target)


def _display_text(cell):
    if cell.value is None:
        return ""
    if isinstance(cell.value, (int, float)) and cell.number_format == "0.0000":
        return f"{cell.value:.4f}"
    return str(cell.value)


def _merged_cell_extents(ws):
    """Index numeric coordinates once per pass; never retain stale worksheet state.

    Construction is linear in the number of merged coordinates (already allocated
    by openpyxl). Lookups are O(1), without parsing range strings for every cell.
    First-range precedence also matches the previous scan for overlapping ranges.
    """
    extents = {}
    for merged in ws.merged_cells.ranges:
        extent = merged.min_col, merged.max_col, merged.min_row, merged.max_row
        for row in range(merged.min_row, merged.max_row + 1):
            for column in range(merged.min_col, merged.max_col + 1):
                extents.setdefault((row, column), extent)
    return extents


def _cell_extent(ws, cell, merged_extents=None):
    if merged_extents is None:
        merged_extents = _merged_cell_extents(ws)
    return merged_extents.get((cell.row, cell.column),
                              (cell.column, cell.column, cell.row, cell.row))


def _needed_height(text, width):
    # Conservative character budget for Calibri 11, including wide glyphs.
    budget = max(1, int((width - 3) / 1.25))
    return 17 * sum(max(1, math.ceil(len(line) / budget)) for line in text.split('\n')) + 6


def apply_plain_presentation(ws):
    """White editable cells, visible gridlines, content-sized columns and rows."""
    merged_extents = _merged_cell_extents(ws)
    for row in ws:
        for cell in row:
            cell.fill = PatternFill(fill_type=None)
            cell.border = Border()
            cell.font = Font(name="Calibri", size=11, color="000000")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for column in range(1, ws.max_column + 1):
        width = 16.0
        for row in ws:
            cell = row[column - 1]
            if cell.value is None:
                continue
            first, last, _, _ = _cell_extent(ws, cell, merged_extents)
            text = _display_text(cell)
            needed = max(map(len, text.split('\n'))) * 1.25 + 4
            if isinstance(cell.value, (int, float)):
                width = max(width, needed)
            elif first == last:
                width = max(width, min(72, needed))
        if width > 255:
            raise ValueError(f"Numeric value exceeds Excel's readable column width: {ws.title}/{column}")
        ws.column_dimensions[get_column_letter(column)].width = width
    for row in range(1, ws.max_row + 1):
        ws.row_dimensions[row].height = 24
    for row in ws:
        for cell in row:
            if cell.value is None:
                continue
            first, last, top, bottom = _cell_extent(ws, cell, merged_extents)
            width = sum(ws.column_dimensions[get_column_letter(c)].width for c in range(first, last + 1))
            height = _needed_height(_display_text(cell), width)
            for r in range(top, bottom + 1):
                ws.row_dimensions[r].height = max(ws.row_dimensions[r].height, height / (bottom - top + 1))
    ws.sheet_view.showGridLines = True
    ws.sheet_view.zoomScale = 100
    ws.sheet_properties.pageSetUpPr.fitToPage = False
    ws.page_setup.fitToWidth = ws.page_setup.fitToHeight = 0
    ws.page_setup.scale = 100
    ws.page_setup.paperSize = None
    ws.print_area = ws.calculate_dimension()


def validate_plain_workbook(wb):
    """Check actual serialized presentation and conservative text extents."""
    for ws in wb:
        merged_extents = _merged_cell_extents(ws)
        if ws.sheet_view.showGridLines is not True or ws.sheet_properties.pageSetUpPr.fitToPage:
            raise ValueError(f"Gridlines/print scaling invalid: {ws.title}")
        if ws.page_setup.fitToWidth or ws.page_setup.fitToHeight or ws.page_setup.paperSize == ws.PAPERSIZE_A3:
            raise ValueError(f"Compressed manuscript print layout: {ws.title}")
        # A full used-range print area is set by apply_plain_presentation.
        from openpyxl.utils.cell import range_boundaries
        if ws.print_area:
            area = str(ws.print_area).split('!')[-1].replace('$', '')
            if range_boundaries(area) != range_boundaries(ws.calculate_dimension()):
                raise ValueError(f"Truncated print area: {ws.title}")
        for row in ws:
            for cell in row:
                if cell.fill.fill_type or any(getattr(cell.border, side).style for side in ('left', 'right', 'top', 'bottom') if getattr(cell.border, side)):
                    raise ValueError(f"Decorative styling: {ws.title}/{cell.coordinate}")
                if cell.value is None:
                    continue
                if cell.font.color is None or cell.font.color.type != 'rgb' or cell.font.color.rgb[-6:] != '000000':
                    raise ValueError(f"Nonblack text: {ws.title}/{cell.coordinate}")
                first, last, top, bottom = _cell_extent(ws, cell, merged_extents)
                width = sum(ws.column_dimensions[get_column_letter(c)].width for c in range(first, last + 1))
                height = sum(ws.row_dimensions[r].height or 15 for r in range(top, bottom + 1))
                if height + .01 < _needed_height(_display_text(cell), width) or not cell.alignment.wrap_text:
                    raise ValueError(f"Potential clipped content: {ws.title}/{cell.coordinate}")
