"""Synthetic reporting checks only: no images, model training or optimizer runs."""
from contextlib import ExitStack, redirect_stdout
import io
from pathlib import Path
import pickle
import tempfile
import unittest
from unittest.mock import patch

import numpy as np
from openpyxl import load_workbook

import main_best as framework


METRICS = {
    "Accuracy": ("AccRuns", True),
    "Precision": ("PSRuns", True),
    "Recall": ("RSRuns", True),
    "F1Score": ("F1Runs", True),
    "Fitness": ("FitRuns", False),
    "Features": ("FeatRuns", False),
    "Time": ("TimeRuns", False),
}


class StatisticalReportingTests(unittest.TestCase):
    def setUp(self):
        self.enterContext(patch.object(framework, 'PLOT_ESTIMATORS', ['knn', 'svm']))
        temporary = tempfile.TemporaryDirectory(prefix="statistical-reporting-")
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.args = framework.parse_args([
            "--optimizers", "DE", "PSO", "--estimators", "knn", "svm",
            "--transfer-functions", "vstf_01", "sstf_02",
            "--output-root", str(self.root), "--reuse-cache-from-exp-id", "none",
        ])

    def workbook(self, results, datasets, args=None):
        args = args or self.args
        path = str(self.root / "statistics.xlsx")
        self.assertEqual(framework.export_statistical_excel(
            results, datasets, args.optimizers, args, path), path)
        workbook = load_workbook(path, data_only=True)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames, list(METRICS))
        return workbook

    def test_known_statistics_filter_nonfinite_and_use_sample_std(self):
        values = [1, np.nan, 3, np.inf, 5, -np.inf]
        self.assertEqual(framework._run_stats(values, "max"),
                         {"Best": 5.0, "Worst": 1.0, "Mean": 3.0, "Std": 2.0})
        self.assertEqual(framework._run_stats(values, "min"),
                         {"Best": 1.0, "Worst": 5.0, "Mean": 3.0, "Std": 2.0})
        for mode in ("min", "max"):
            self.assertEqual(framework._run_stats([np.nan, 7, np.inf], mode),
                             {"Best": 7.0, "Worst": 7.0, "Mean": 7.0, "Std": 0.0})
            for missing in ([], [np.nan, np.inf, -np.inf]):
                self.assertTrue(all(np.isnan(value) for value in framework._run_stats(missing, mode).values()))

    def test_workbook_preserves_every_combination_and_metric_direction(self):
        results = {}
        expected = {}
        headers = []
        for dataset_index, dataset in enumerate(["First", "Second"]):
            results[dataset] = {}
            for method_index, method in enumerate(self.args.optimizers):
                for classifier_index, classifier in enumerate(self.args.estimators):
                    for transfer_index, transfer in enumerate(self.args.transfer_functions):
                        header = f"{method} | {classifier.upper()} | {transfer.upper()}"
                        if dataset_index == 0:
                            headers.append(header)
                        label = framework.build_alg_label(method, transfer, classifier, True, True)
                        row = {"Estimator": classifier, "AccMean": -999.0}
                        for metric_index, (sheet, (key, maximize)) in enumerate(METRICS.items()):
                            offset = 1000 * dataset_index + 100 * method_index + 10 * classifier_index + transfer_index + metric_index
                            row[key] = np.array([offset + 1, np.nan, offset + 3, np.inf, offset + 5, -np.inf])
                            expected[dataset, header, sheet] = (
                                [offset + 5, offset + 1, offset + 3, 2.0] if maximize else
                                [offset + 1, offset + 5, offset + 3, 2.0]
                            )
                        results[dataset][label] = row
        original = pickle.dumps(results)
        book = self.workbook(results, ["First", "Second"])
        self.assertEqual(pickle.dumps(results), original)
        for sheet in METRICS:
            ws = book[sheet]
            self.assertEqual([cell.value for cell in ws[1]], ["Dataset", "Statistic"] + headers)
            self.assertEqual({str(merged) for merged in ws.merged_cells.ranges}, {"A2:A5", "A6:A9"})
            for dataset_index, dataset in enumerate(["First", "Second"]):
                first_row = 2 + 4 * dataset_index
                self.assertEqual(ws.cell(first_row, 1).value, dataset)
                self.assertEqual([ws.cell(first_row + i, 2).value for i in range(4)],
                                 ["Best", "Worst", "Mean", "Std"])
                for column, header in enumerate(headers, start=3):
                    self.assertEqual([ws.cell(first_row + i, column).value for i in range(4)],
                                     expected[dataset, header, sheet])

    def test_single_configuration_legacy_label_and_missing_metrics(self):
        self.args.optimizers = ["OriginalPSO"]
        self.args.estimators = ["knn"]
        self.args.transfer_functions = ["vstf_01"]
        book = self.workbook({"Single": {"ORIGINALPSO": {
            "AccRuns": [np.nan, 87.5, np.inf], "FitRuns": [np.nan, -np.inf],
        }}}, ["Single", "Absent"])
        for sheet in METRICS:
            ws = book[sheet]
            self.assertEqual([cell.value for cell in ws[1]],
                             ["Dataset", "Statistic", "PSO | KNN | VSTF_01"])
            expected = [87.5, 87.5, 87.5, 0] if sheet == "Accuracy" else [None] * 4
            self.assertEqual([ws.cell(row, 3).value for row in range(2, 6)], expected)
            self.assertEqual([ws.cell(row, 3).value for row in range(6, 10)], [None] * 4)

    def test_available_combinations_outside_configuration_are_retained(self):
        self.args.optimizers = ["DE"]
        self.args.estimators = ["knn"]
        self.args.transfer_functions = ["vstf_01"]
        book = self.workbook({"Data": {
            "DE": {"Estimator": "knn", "AccRuns": [80]},
            "MACRO-DE-T_SSTF_02_SVM": {"Estimator": "svm", "AccRuns": [90]},
        }}, ["Data"])
        self.assertEqual([cell.value for cell in book["Accuracy"][1]], [
            "Dataset", "Statistic", "DE | KNN | VSTF_01", "MaCRO-DE-t | SVM | SSTF_02",
        ])
        self.assertEqual(book["Accuracy"].cell(2, 4).value, 90)

    def test_main_figures_only_from_legacy_and_partial_caches_never_runs_science(self):
        args = framework.parse_args([
            "--figures-only", "--dataset-name", "Synthetic", "--pipeline-mode", "full",
            "--miafex-dataset-root", str(self.root / "absent-images"),
            "--feature-dataset-root", str(self.root / "absent-features"),
            "--output-root", str(self.root), "--reuse-cache-from-exp-id", "none",
            "--optimizers", "DE", "--estimators", "knn", "svm", "--runs", "20",
            "--transfer-functions", "vstf_01",
        ])
        args.optimizers = framework.resolve_optimizers(args)
        scoped = framework.resolve_miafex_dataset_args(args)["Synthetic"]
        signature = framework.build_cache_signature(scoped)
        paths = framework.make_paths(args)
        snapshots = {}
        expected_results = {"Synthetic": {}}
        for classifier, run_ids in (("knn", None), ("svm", [2, 8, 15])):
            values = [1.0, 3.0, 5.0] if classifier == "knn" else [11.0, 13.0, 15.0]
            row = framework.build_label_payload(
                classifier, *[values for _ in METRICS],
                [np.full(args.epochs, value) for value in values], args.epochs, completed_run_ids=run_ids,
            )
            # Both classifier caches deliberately use the same legacy label.
            label = framework.build_alg_label(args.optimizers[0], "vstf_01", classifier, False, False)
            expected_results["Synthetic"][label] = row
            # One final legacy payload and one partial, non-contiguous progress payload.
            filename = Path(framework.cache_files(paths, "Synthetic", classifier, signature)[int(run_ids is not None)])
            framework.save_cache(str(filename), {label: row})
            snapshots[filename] = (filename.read_bytes(), filename.stat().st_mtime_ns)
        expected_summary = framework.generate_summary_dataframe(expected_results, args)
        with ExitStack() as stack:
            stack.enter_context(patch.object(framework, "parse_args", return_value=args))
            stack.enter_context(patch.object(framework, "resolve_execution_config"))
            stack.enter_context(patch.object(framework, "print_backend_report"))
            stack.enter_context(patch.object(framework, "validate_execution_config"))
            chart = stack.enter_context(patch.object(framework, "generate_seven_global_charts", return_value=[]))
            for name in ("train_miafex", "extract_miafex_features", "run_single", "execute_pending_runs",
                         "resolve_miafex_csv", "load_miafex_feature_data", "get_dataset", "save_cache"):
                stack.enter_context(patch.object(framework, name, side_effect=AssertionError(f"Forbidden: {name}")))
            with redirect_stdout(io.StringIO()):
                framework.main()
        for filename, snapshot in snapshots.items():
            self.assertEqual((filename.read_bytes(), filename.stat().st_mtime_ns), snapshot)
        self.assertEqual(framework.build_cache_signature(scoped), signature)
        framework.pd.testing.assert_frame_equal(chart.call_args.args[0], expected_summary)
        csv_path = Path(paths.res_dir) / f"RESUMEN_GRAFICAS_{paths.exp_tag}.csv"
        self.assertEqual(csv_path.read_text(), expected_summary.to_csv(index=False))
        with (Path(paths.res_dir) / f"Statistical_Results_{paths.exp_tag}.xlsx").open("rb") as stream:
            book = load_workbook(stream, data_only=True)
            self.assertEqual(book.sheetnames, list(METRICS))
            self.assertEqual([cell.value for cell in book["Accuracy"][1]], [
                "Dataset", "Statistic", "DE | KNN | VSTF_01", "DE | SVM | VSTF_01",
            ])
            for column in (3, 4):
                expected = [5, 1, 3, 2] if column == 3 else [15, 11, 13, 2]
                self.assertEqual([book["Accuracy"].cell(row, column).value for row in range(2, 6)], expected)
            book.close()

    def test_normal_reporting_from_completed_cache_adds_workbook(self):
        args = framework.parse_args([
            "--dataset-name", "Synthetic", "--dataset-root", str(self.root / "absent-images"),
            "--pipeline-mode", "feature_selection", "--feature-dataset-root", str(self.root / "features"),
            "--output-root", str(self.root), "--reuse-cache-from-exp-id", "none",
            "--optimizers", "DE", "--estimators", "knn", "--transfer-functions", "vstf_01", "--runs", "3",
        ])
        args.optimizers = framework.resolve_optimizers(args)
        scoped = framework.resolve_miafex_dataset_args(args)["Synthetic"]
        for filename in framework.miafex_feature_paths(scoped).values():
            path = Path(filename)
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text("f0,label\n1,0\n2,1\n")
        paths = framework.make_paths(args)
        signature = framework.build_cache_signature(scoped)
        row = framework.build_label_payload("knn", *[[1, 3, 5] for _ in METRICS],
                                            [np.full(args.epochs, value) for value in (1, 3, 5)],
                                            args.epochs, completed_run_ids=[0, 1, 2])
        label = framework.build_alg_label(args.optimizers[0], "vstf_01", "knn", False, False)
        framework.save_cache(framework.cache_files(paths, "Synthetic", "knn", signature)[0], {label: row})
        with ExitStack() as stack:
            stack.enter_context(patch.object(framework, "parse_args", return_value=args))
            stack.enter_context(patch.object(framework, "resolve_execution_config"))
            stack.enter_context(patch.object(framework, "print_backend_report"))
            stack.enter_context(patch.object(framework, "validate_execution_config"))
            stack.enter_context(patch.object(framework, "generate_seven_global_charts", return_value=[]))
            for name in ("train_miafex", "extract_miafex_features", "run_single", "execute_pending_runs"):
                stack.enter_context(patch.object(framework, name, side_effect=AssertionError(f"Forbidden: {name}")))
            with redirect_stdout(io.StringIO()):
                framework.main()
        for prefix in ("Global_Results", "Statistical_Results"):
            with (Path(paths.res_dir) / f"{prefix}_{paths.exp_tag}.xlsx").open("rb") as stream:
                book = load_workbook(stream, data_only=True)
                self.assertEqual(book.sheetnames, list(METRICS))
                if prefix == "Global_Results":
                    self.assertEqual(book["Accuracy"].cell(2, 2).value, 3)
                else:
                    self.assertEqual([book["Accuracy"].cell(row, 3).value for row in range(2, 6)], [5, 1, 3, 2])
                book.close()

    def test_empty_results_still_generate_seven_sheets(self):
        book = self.workbook({}, [])
        headers = [f"{method} | {classifier.upper()} | {tf.upper()}"
                   for method in self.args.optimizers for classifier in self.args.estimators
                   for tf in self.args.transfer_functions]
        for sheet in METRICS:
            self.assertEqual([cell.value for cell in book[sheet][1]], ["Dataset", "Statistic"] + headers)


if __name__ == "__main__":
    unittest.main()
