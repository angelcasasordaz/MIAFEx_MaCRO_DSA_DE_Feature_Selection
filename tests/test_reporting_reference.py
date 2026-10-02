"""Read-only EXP604 regression against validated MIAFEx scientific outputs."""
from contextlib import ExitStack, redirect_stdout
from io import StringIO
from pathlib import Path
import json
import sys
import tempfile
import unittest
from unittest.mock import patch

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.stats import friedmanchisquare, rankdata, t

import main_best as m
from reporting import core, figures, statistics, paper_tables


def exact_signed_rank_check(differences):
    """Independent subset-sum oracle for nonzero, untied signed ranks."""
    d = np.asarray(differences)
    assert np.all(d != 0) and len(np.unique(abs(d))) == len(d)
    ranks = rankdata(abs(d)).astype(int)
    positive, total = int(ranks[d > 0].sum()), int(ranks.sum())
    counts = np.zeros(total + 1, dtype=np.int64)
    counts[0] = 1
    for rank in ranks:
        counts[rank:] += counts[:-rank].copy()
    w = min(positive, total - positive)
    return w, min(1., 2 * counts[:w+1].sum() / (2**len(d)))


class HistoricalScientificValuesTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.root = Path(__file__).resolve().parents[1]
        cls.source = cls.root / 'Results/EXP604'
        if not (cls.source / 'Global_Results_EXP604.xlsx').is_file():
            raise unittest.SkipTest('Saved EXP604 regression outputs unavailable')
        args = m.parse_args(['--report-only', '--exp-id', '604', '--output-root', str(cls.root)])
        # Historical scientific configuration is confined to this fixture.
        args.miafex_datasets = ['Brain_MRI', 'Breast_Ultrasound', 'Chest_CT', 'Eye_Fundus',
                               'Gastrointestinal_Endoscopy', 'Histological_Biopsy', 'Ocular_Alignment']
        args.optimizers = ['MaCRO-DE-t', 'DE', 'JADE', 'SHADE', 'PSO', 'GWO', 'WOA',
                           'HHO', 'BRO', 'DBO', 'RUN', 'FOX', 'FLA']
        args.estimators, args.transfer_functions = ['knn', 'svm'], ['vstf_01']
        args.runs, args.epochs, args.pop_size, args.seed_base = 20, 200, 30, 1234
        args.random_state, args.test_size, args.dsade_mahal_q = 42, .2, .68
        with core.report_guard(()):
            cls.report = core.load_completed_cache(args)
        cls.hashes = {p: (core.sha256(p), p.stat().st_mtime_ns) for kind in ('Results', 'Figures')
                      for p in (cls.root / kind / 'EXP604').rglob('*') if p.is_file()}

    @classmethod
    def tearDownClass(cls):
        for path, (digest, mtime) in cls.hashes.items():
            if core.sha256(path) != digest or path.stat().st_mtime_ns != mtime:
                raise AssertionError(f'Historical cache/report changed: {path}')

    def setUp(self):
        stack = ExitStack(); self.addCleanup(stack.close)
        for name in ('run_single', 'execute_pending_runs', 'build_optimizer', 'save_cache',
                     'save_combination', 'get_dataset', 'train_miafex', 'extract_miafex_features',
                     'resolve_miafex_csv', 'load_miafex_feature_data', 'resolve_execution_config'):
            blocked = stack.enter_context(patch.object(m, name, side_effect=AssertionError('Scientific execution forbidden')))
            self.addCleanup(blocked.assert_not_called)

    def test_frozen_sources_and_historical_science_are_preserved(self):
        frozen = json.loads((self.root / 'diagnostics/exp604_cache_migration/manifest.json').read_text())
        self.assertEqual(len(self.report.indexed), 182)
        self.assertEqual(sum(row['CompletedRuns'] for row in self.report.indexed.values()), 3640)
        for filename, expected in frozen['original_files'].items():
            path = self.source / 'cache' / filename
            self.assertEqual(core.sha256(path), expected['sha256'])
            self.assertEqual(path.stat().st_mtime_ns, expected['mtime_ns'])
        self.assertEqual({i['wrapper_revision'] for i in self.report.identities.values()},
                         {'mafese-prepared-partitions-default-selector-v1'})
        self.assertNotEqual(next(iter(self.report.identities.values()))['wrapper_revision'], m.WRAPPER_SCIENCE_REVISION)

    def test_generic_means_ci_and_radar_preserve_historical_values(self):
        self.enterContext(patch.object(m, 'PLOT_RUN_AGGREGATION', 'mean'))
        reference = pd.read_csv(self.source / 'RESUMEN_GRAFICAS_EXP604.csv').rename(
            columns={'Archivo': 'Dataset', 'Estimador': 'Estimator', 'Optimizador': 'Optimizer'})
        reference['Optimizer'] = reference.Optimizer.map(m.optimizer_acronym)
        report = self.report
        columns = {'AccRuns': 'AS_test', 'PSRuns': 'PS_test', 'RSRuns': 'RS_test',
                   'F1Runs': 'F1_test', 'FeatRuns': 'N_Features_Selected', 'TimeRuns': 'Runtime'}
        for classifier in report.classifiers:
            for metric in report.metrics:
                if metric.run_key in columns:
                    expected = figures.metric_values(reference, columns[metric.run_key], classifier,
                                                     report.datasets, report.algorithms)
                    np.testing.assert_allclose(figures.metric_matrix(report, classifier, metric), expected,
                                               rtol=1e-12, atol=1e-12)
            precision = next(metric for metric in report.metrics if metric.run_key == 'PSRuns')
            values = figures.metric_matrix(report, classifier, precision)
            means, ci = figures.dataset_mean_ci(values)
            np.testing.assert_allclose(means, values.mean(axis=1), rtol=1e-14)
            np.testing.assert_allclose(ci, t.ppf(.975, 6) * values.std(axis=1, ddof=1) / np.sqrt(7), rtol=1e-14)
            labels, radar = figures.radar_values(report, classifier)
            self.assertEqual(len(labels), 5)
            features = figures.metric_values(reference, 'N_Features_Selected', classifier,
                                              report.datasets, report.algorithms)
            np.testing.assert_allclose(radar[:, :, -1], 1 - features / np.maximum(features.max(axis=0), 1), rtol=1e-12)

    def test_statistics_match_independent_matched_blocks_and_signed_ranks(self):
        report = self.report
        metric = next(metric for metric in report.metrics if metric.run_key == 'F1Runs')
        blocks, x = statistics.matched_block_matrix(report.indexed, report.datasets, report.classifiers,
                                                    report.algorithms, metric)
        expected = np.array([[report.indexed[ds, cls, opt]['F1Runs'].mean() for opt in report.algorithms]
                             for ds, cls in blocks])
        np.testing.assert_allclose(x, expected)
        analysis = statistics.analyze(x, report.algorithms)
        np.testing.assert_allclose(analysis['mean_ranks'], rankdata(-expected, axis=1).mean(axis=0))
        self.assertAlmostEqual(analysis['friedman']['statistic'], friedmanchisquare(*expected.T).statistic)
        self.assertEqual(len(blocks), 14)
        self.assertEqual(len(analysis['pairs']), 78)
        for row in analysis['pairs'].itertuples():
            if row.Method == 'exact':
                i, j = report.algorithms.index(row.Algorithm_A), report.algorithms.index(row.Algorithm_B)
                w, p = exact_signed_rank_check(x[:, i] - x[:, j])
                self.assertEqual(row.Wilcoxon_statistic, w)
                self.assertAlmostEqual(row.Raw_p, p, places=14)

    def test_exp604_cli_exports_in_temporary_destination_and_matches_saved_excel(self):
        from tests.test_generic_reporting import tiny_figures
        report = self.report
        with tempfile.TemporaryDirectory() as folder:
            argv = ['main_best.py', '--report-only', '--exp-id', '604', '--output-root', str(self.root),
                    '--report-output-root', folder, '--miafex-datasets', *report.datasets,
                    '--optimizers', *report.args.optimizers, '--estimators', *report.classifiers]
            # Actual dispatch, Excel/statistics and validation; tiny figures keep this lightweight.
            with patch.object(sys, 'argv', argv), \
                    patch.object(figures, 'publication_figures', tiny_figures), \
                    patch.object(figures, 'statistical_figures', tiny_figures), redirect_stdout(StringIO()):
                manifest = m.main()
            self.assertEqual(manifest['experiment_id'], 604)
            self.assertEqual(manifest['optimization_calls'], 0)
            self.assertTrue(manifest['protected_files_unchanged'])
            root = Path(folder)
            res = root / 'Results/EXP604/full_rep1'
            self.assertEqual(json.loads((root / 'Figures/EXP604/full_rep1/validation.json').read_text()), manifest)
            self.assertFalse(list(root.rglob('*.pkl')))
            self.assertFalse(list(root.rglob('*.pdf')))
            for name in ('Global_Results', 'Statistical_Results', 'Full_Friedman_Analysis'):
                filename = f'{name}_EXP604.xlsx'
                actual = pd.read_excel(res / filename, sheet_name=None)
                reference = pd.read_excel(self.source / filename, sheet_name=None)
                self.assertEqual(list(actual), list(reference))
                for sheet in reference:
                    pd.testing.assert_frame_equal(actual[sheet], reference[sheet], check_dtype=False,
                                                  check_exact=False, rtol=1e-12, atol=1e-12)
            wb = load_workbook(res / 'Paper_Tables_EXP604.xlsx', data_only=True)
            try:
                paper_tables.validate_plain_workbook(wb)
                for ai, opt in enumerate(report.algorithms):
                    for mi, metric in enumerate(report.metrics):
                        for ci, cls in enumerate(report.classifiers):
                            vectors = [report.indexed[ds, cls, opt][metric.run_key] / metric.scale for ds in report.datasets]
                            vector = np.mean(vectors, axis=0)
                            expected = [max(vector), min(vector), np.mean(vector), np.std(vector, ddof=1)]
                            if metric.best_mode == 'min': expected[:2] = expected[1::-1]
                            actual = [wb['Overall'].cell(3+4*ai+s, 3+2*mi+ci).value for s in range(4)]
                            np.testing.assert_allclose(actual, expected, rtol=1e-12, atol=1e-12)
            finally:
                wb.close()
