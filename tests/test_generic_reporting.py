"""Small synthetic completed caches; mocked scientific paths and tiny PNGs."""
from contextlib import ExitStack
from copy import deepcopy
from io import StringIO
from pathlib import Path
import json
import pickle
import sys
import tempfile
import unittest
from unittest.mock import patch

import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook

import main_best as m
from reporting import core, figures, statistics, paper_tables

ORIGINAL_FEATURE_LOADER = m.load_miafex_feature_data

def arguments(root):
    return m.parse_args(['--report-only', '--exp-id', '913', '--output-root', str(root),
                         '--miafex-datasets', 'First', 'Second',
                         '--feature-dataset-root', str(root / 'features'),
                         '--optimizers', 'DE', 'PSO', 'JADE', '--estimators', 'knn', 'rf',
                         '--runs', '3', '--epochs', '3', '--pop-size', '5'])


def caches(args):
    """Real MIAFEx v2 identities and synthetic completed rows; no model fitting."""
    import scientific_cache
    local = deepcopy(args)
    local.optimizers = m.resolve_optimizers(local)
    for ds in local.miafex_datasets:
        directory = Path(local.feature_dataset_root) / ds
        directory.mkdir(parents=True, exist_ok=True)
        for split in ('train', 'test'):
            (directory / f'{split}_features.csv').write_text('f0,f1,label\n1,2,0\n3,4,1\n')
    scoped = m.resolve_miafex_dataset_args(local)
    paths = m.make_paths(local, create=False)
    for di, ds in enumerate(local.miafex_datasets):
        for ci, cls in enumerate(local.estimators):
            for ai, method in enumerate(local.optimizers):
                for ti, transfer in enumerate(local.transfer_functions):
                    identity = m.build_combination_identity(scoped[ds], ds, cls, method, transfer)
                    base = np.array([.2, .3, .4]) + di * .01 + ci * .005 + ai * .002 + ti * .001
                    row = {'Estimator': cls, 'CompletedRuns': local.runs,
                           'CompletedRunIDs': list(range(local.runs)),
                           'Curve': [.9, .7, float(base.mean())],
                           'CurvesAll': [[.9, .7, float(v)] for v in base]}
                    for metric in paper_tables.METRICS:
                        values = base * metric.scale
                        row[metric.run_key] = values.tolist()
                        row[metric.run_key.replace('Runs', 'Mean')] = float(values.mean())
                    path, _ = scientific_cache.combination_files(paths, identity)
                    path.parent.mkdir(parents=True, exist_ok=True)
                    path.write_bytes(pickle.dumps({'identity': identity, 'row': row}))
    return args


def tiny_figures(*args):
    fig = plt.figure(figsize=(.5, .5))
    yield 'tiny', fig


class GenericReportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.stack = ExitStack(); self.addCleanup(self.stack.close)
        for name in ('run_single', 'execute_pending_runs', 'build_optimizer', 'get_dataset', 'save_cache',
                     'resolve_execution_config', 'resolve_miafex_csv', 'train_miafex', 'extract_miafex_features',
                     'load_miafex_feature_data', 'resolve_cached_payload'):
            mock = self.stack.enter_context(patch.object(m, name, side_effect=AssertionError('Optimization forbidden')))
            if name == 'load_miafex_feature_data':
                mock.__wrapped__ = ORIGINAL_FEATURE_LOADER  # Identity introspection only; calls remain forbidden.
            self.addCleanup(mock.assert_not_called)

    def test_f_missing_or_incomplete_cache_fails_without_allocation(self):
        args = arguments(self.root)
        with self.assertRaises(FileNotFoundError): core.run_report(args)
        self.assertFalse((self.root / 'Figures').exists())
        caches(args)
        path = next(self.root.rglob('*.pkl'))
        with path.open('rb') as stream: payload = pickle.load(stream)
        payload['row']['CompletedRuns'] = 2
        with path.open('wb') as stream: pickle.dump(payload, stream)
        with self.assertRaisesRegex(ValueError, 'Incomplete'): core.run_report(args)
        self.assertFalse((self.root / 'Figures').exists())

    def test_separate_report_destination_reads_original_caches(self):
        args = arguments(self.root)
        caches(args)
        args.report_output_root = str(self.root / 'generated')
        before = {p: core.sha256(p) for p in self.root.rglob('*.pkl')}
        with patch.object(figures, 'publication_figures', tiny_figures), patch.object(figures, 'statistical_figures', tiny_figures):
            manifest = core.run_report(args)
        self.assertEqual(manifest['optimization_calls'], 0)
        self.assertFalse((self.root / 'Figures').exists())
        destination = self.root / 'generated'
        self.assertTrue((destination / 'Results/EXP913/full_rep1/Paper_Tables_EXP913.xlsx').is_file())
        self.assertFalse(list(destination.rglob('*.pkl')))
        self.assertTrue(all(core.sha256(path) == digest for path, digest in before.items()))

    def test_g_complete_outputs_versions_and_unchanged_sources(self):
        args = arguments(self.root); caches(args)
        before_args = deepcopy(vars(args))
        before = {p: core.sha256(p) for p in self.root.rglob('*.pkl')}
        with patch.object(figures, 'publication_figures', tiny_figures), patch.object(figures, 'statistical_figures', tiny_figures):
            first = core.run_report(args)
            previous = {p: core.sha256(p) for p in self.root.rglob('*') if p.is_file()}
            second = core.run_report(args)
        self.assertEqual(vars(args), before_args)
        self.assertEqual((first['report_version'], second['report_version']), (1, 2))
        self.assertEqual(second['experiment_id'], 913)
        self.assertEqual(second['optimization_calls'], 0)
        for paths in (before, previous):
            self.assertTrue(all(core.sha256(p) == digest for p, digest in paths.items()))
        res = self.root / 'Results/EXP913/full_rep2'
        expected = {'Global_Results_EXP913.xlsx', 'Statistical_Results_EXP913.xlsx',
                    'Full_Friedman_Analysis_EXP913.xlsx', 'Paper_Tables_EXP913.xlsx'}
        self.assertEqual({p.name for p in res.glob('*.xlsx')}, expected)
        self.assertTrue((res / 'statistics/statistical_report.txt').is_file())
        self.assertTrue((res / 'statistics/matched_block_means.csv').is_file())
        self.assertEqual(json.loads((res / 'validation.json').read_text())['pdfs_generated'], 0)
        self.assertFalse(list(self.root.rglob('*.pdf')))
        self.assertFalse((self.root / 'Results/EXP914').exists())
        wb = load_workbook(res / 'Paper_Tables_EXP913.xlsx'); self.addCleanup(wb.close)
        paper_tables.validate_plain_workbook(wb)
        # Existing exporters have precisely the same cells as independent exports.
        local = deepcopy(args); local.experiment_mode = 'full'
        report = core.load_completed_cache(local)
        for filename, exporter, params in (
            ('Global_Results_EXP913.xlsx', m.export_global_excel, (report.results, report.datasets)),
            ('Statistical_Results_EXP913.xlsx', m.export_statistical_excel, (report.results, report.datasets, args.optimizers, local)),
            ('Full_Friedman_Analysis_EXP913.xlsx', m.export_friedman_analysis, (report.results, report.datasets, args.optimizers, local)),
        ):
            direct = self.root / 'direct.xlsx'
            exporter(*params, str(direct))
            a, b = load_workbook(res / filename), load_workbook(direct)
            try:
                self.assertEqual(a.sheetnames, b.sheetnames)
                for sheet in a.sheetnames: self.assertEqual(list(a[sheet].values), list(b[sheet].values))
            finally: a.close(); b.close()

    def test_failure_during_export_removes_only_new_staging(self):
        args = arguments(self.root); caches(args)
        with patch.object(core, 'generate_outputs', side_effect=ValueError('failed workbook')):
            with self.assertRaisesRegex(ValueError, 'failed workbook'): core.run_report(args)
        self.assertEqual(core.next_report_version(self.root, 913), 1)
        self.assertFalse(list(self.root.rglob('.report-staging-*')))
        self.assertEqual(len(list(self.root.rglob('*.pkl'))), 12)

    def test_h_variable_plot_dimensions_and_missing_curves(self):
        for datasets, algorithms, classifiers in [(['One'], ['DE'], ['tree']),
                                                   (['A', 'B', 'C', 'D'], ['DE', 'PSO'], ['knn', 'rf', 'svm'])]:
            from tests.test_paper_tables import fixture
            args, results = fixture(datasets, algorithms, classifiers)
            indexed = {(d, r['Estimator'], label.rsplit('_', 1)[0]): r for d, rows in results.items() for label, r in rows.items()}
            report = core.CompletedReport(args, results, indexed, datasets, classifiers, algorithms,
                                          list(paper_tables.METRICS[:1]), 'fixture', {})
            skipped = []
            generated = list(figures.publication_figures(report, skipped))
            try:
                self.assertEqual(len(generated), 3 * len(classifiers))
                self.assertEqual(sum(item['output'].startswith('Convergence') for item in skipped),
                                 len(datasets) * len(classifiers))
                for _, fig in generated:
                    labels = [label.get_text() for ax in fig.axes for label in ax.get_xticklabels()+ax.get_yticklabels()]
                    self.assertTrue(set(algorithms).issubset(labels))
            finally:
                for _, fig in generated: plt.close(fig)

    def test_plain_presentation_long_names_full_range_and_preserved_numbers(self):
        from tests.test_paper_tables import fixture
        args, results = fixture(datasets=('Long dataset name with a multiline\nheading and more text',), classifiers=('longclassifier',))
        indexed = {(d, r['Estimator'], label.rsplit('_', 1)[0]): r for d, rows in results.items() for label, r in rows.items()}
        path = self.root / 'Paper_Tables_fixture.xlsx'
        paper_tables.export_indexed_tables(indexed, list(results), args.optimizers, args.estimators,
                                          list(paper_tables.METRICS), path, title='Fixture')
        wb = load_workbook(path); self.addCleanup(wb.close)
        paper_tables.validate_plain_workbook(wb)
        for ws in wb:
            self.assertTrue(ws.sheet_view.showGridLines)
            self.assertGreaterEqual(ws.row_dimensions[1].height, 24)
            self.assertTrue(all(dim.width >= 16 for dim in ws.column_dimensions.values()))
            self.assertEqual(ws.page_setup.fitToWidth, 0)
            self.assertEqual(ws.page_setup.fitToHeight, 0)


class StatisticalTests(unittest.TestCase):
    def test_i_variable_blocks_friedman_and_holm(self):
        from scipy.stats import friedmanchisquare, rankdata
        for n, k in ((3, 3), (7, 4), (2, 2), (1, 1)):
            x = np.arange(n*k, dtype=float).reshape(n, k) ** 1.3
            result = statistics.analyze(x, [str(i) for i in range(k)])
            np.testing.assert_allclose(result['ranks'], rankdata(-x, axis=1))
            self.assertEqual(len(result['pairs']), k*(k-1)//2)
            if n >= 2 and k >= 3:
                self.assertAlmostEqual(result['friedman']['statistic'], friedmanchisquare(*x.T).statistic)
            raw = result['pairs'].Raw_p.to_numpy()
            order = np.argsort(raw)
            adjusted = result['pairs'].Holm_adjusted_p.to_numpy()
            for rank, i in enumerate(order):
                self.assertAlmostEqual(adjusted[i], min(1, max((len(raw)-j)*raw[order[j]] for j in range(rank+1))))

    def test_missing_blocks_zero_differences_ties_and_insufficient_data(self):
        r = statistics.analyze([[1,1,2], [2,2,3], [np.nan,3,4]], ['A','B','C'])
        self.assertEqual(r['omitted_blocks'], [2])
        self.assertEqual(r['pairs'].iloc[0].Raw_p, 1)
        self.assertIn('degenerate', r['pairs'].iloc[0].Method)
        self.assertEqual(r['pairs'].iloc[1].Method, 'approx')
        r = statistics.analyze([[1,1,1], [2,2,2]], ['A','B','C'])
        self.assertIsNone(r['friedman']['statistic'])
        self.assertIn('tied', r['friedman']['status'])
        r = statistics.analyze([[1,2,3]], ['A','B','C'])
        self.assertTrue(r['pairs'].Raw_p.isna().all())
        self.assertTrue(r['summary'].Std.isna().all())

    def test_exact_distribution_independent_fixture(self):
        from tests.test_reporting_reference import exact_signed_rank_check
        x = np.array([[1, 2], [4, 1], [3, 7], [8, 3]], dtype=float)
        result = statistics.analyze(x, ['A','B'])
        w, p = exact_signed_rank_check(x[:,0]-x[:,1])
        self.assertEqual(result['pairs'].iloc[0].Method, 'exact')
        self.assertEqual(result['pairs'].iloc[0].Wilcoxon_statistic, w)
        self.assertEqual(result['pairs'].iloc[0].Raw_p, p)

class AdditionalCacheTests(unittest.TestCase):
    setUp = GenericReportTests.setUp
    def test_available_metrics_and_absent_curves_are_reported(self):
        args = arguments(self.root); caches(args)
        for path in self.root.rglob('*.pkl'):
            with path.open('rb') as stream: payload = pickle.load(stream)
            for row in [payload['row']]:
                for metric in paper_tables.METRICS[1:]:
                    row.pop(metric.run_key); row.pop(metric.run_key.replace('Runs', 'Mean'))
                row.pop('Curve'); row.pop('CurvesAll')
            with path.open('wb') as stream: pickle.dump(payload, stream)
        local = deepcopy(args); local.experiment_mode = 'full'
        report = core.load_completed_cache(local)
        self.assertEqual([m.run_key for m in report.metrics], ['AccRuns'])
        with patch.object(figures, 'publication_figures', tiny_figures), patch.object(figures, 'statistical_figures', tiny_figures):
            manifest = core.run_report(args)
        self.assertTrue(any('FitRuns' in item['reason'] for item in manifest['reports'][0]['skipped_outputs']))
        self.assertFalse(list((self.root / 'Results/EXP913/full_rep1').glob('*Friedman*.xlsx')))

    def test_missing_entire_classifier_cache_and_partial_metric_fail(self):
        args = arguments(self.root); caches(args)
        path = next(self.root.rglob('*.pkl'))
        original = path.read_bytes()
        path.unlink()
        with self.assertRaises(FileNotFoundError): core.run_report(args)
        path.write_bytes(original)
        payload = pickle.loads(original)
        payload['row'].pop('F1Runs')
        with path.open('wb') as stream: pickle.dump(payload, stream)
        with self.assertRaisesRegex(ValueError, 'Incomplete metric grid'): core.run_report(args)
        self.assertFalse((self.root / 'Figures').exists())

    def test_multiple_transfer_functions_keep_distinct_identity(self):
        args = arguments(self.root)
        args.transfer_functions = ['vstf_01', 'sstf_02']
        caches(args)
        report = core.load_completed_cache(args)
        self.assertEqual(len(report.algorithms), 6)
        self.assertEqual(len(report.indexed), 24)
        self.assertTrue(all('_VSTF_01' in opt or '_SSTF_02' in opt for opt in report.algorithms))
        with patch.object(figures, 'publication_figures', tiny_figures), patch.object(figures, 'statistical_figures', tiny_figures):
            manifest = core.run_report(args)
        self.assertEqual(manifest['optimization_calls'], 0)

    def test_single_classifier_labels_remain_qualified(self):
        args = arguments(self.root); args.estimators = ['knn']; caches(args)
        report = core.load_completed_cache(args)
        self.assertEqual(report.classifiers, ['knn'])
        self.assertTrue(all(label.endswith('_KNN') for rows in report.results.values() for label in rows))

    def test_mafese_reporting_uses_cache_metadata_without_loading_datasets(self):
        import scientific_cache
        args = arguments(self.root); caches(args)
        for path in self.root.rglob('*.pkl'):
            entry = pickle.loads(path.read_bytes())
            identity = entry['identity']
            identity['dataset_source'] = 'mafese'
            identity['partition']['mode'] = 'mafese_split_v1'
            identity['prepared_features'] = {'X': 'stored-features', 'y': 'stored-targets'}
            destination = path.with_name(scientific_cache.identity_digest(identity) + '_results.pkl')
            destination.write_bytes(pickle.dumps(entry)); path.unlink()
        args.dataset_source = 'mafese'
        with patch.object(m, 'resolve_mafese_dataset_names', return_value=['First', 'Second']), core.report_guard(()):
            report = core.load_completed_cache(args)
        self.assertEqual(len(report.indexed), 12)

    def test_migrated_references_require_untouched_final_sources(self):
        args = arguments(self.root); caches(args)
        path = next(self.root.rglob('*.pkl'))
        entry = pickle.loads(path.read_bytes())
        identity = entry['identity']
        label = m.build_alg_label(identity['optimizer'], identity['transfer_function'], identity['classifier'], False, True)
        original = path.parent.parent / 'legacy_results.pkl'
        original.write_bytes(pickle.dumps({label: entry['row']}))
        reference = {'identity': identity, 'sources': [
            {'file': original.name, 'label': label, 'sha256': core.sha256(original)},
            {'file': 'missing_progress.pkl', 'label': label, 'sha256': 'unused'}]}
        path.unlink(); path.with_suffix('.json').write_text(json.dumps(reference))
        with core.report_guard(()):
            report = core.load_completed_cache(args)
        self.assertEqual(len(report.indexed), 12)
        self.assertTrue(any(name.endswith('legacy_results.pkl') for name in report.sources))
        original.write_bytes(original.read_bytes() + b'changed')
        with self.assertRaisesRegex(ValueError, 'Original payload changed'):
            core.run_report(args)
        self.assertFalse((self.root / 'Figures').exists())

    def test_progress_and_cross_exp_fallback_never_supply_final(self):
        args = arguments(self.root); caches(args)
        path = next(self.root.rglob('*.pkl'))
        path.rename(path.with_name(path.name.replace('_results', '_progress')))
        with self.assertRaises(FileNotFoundError): core.run_report(args)
        self.assertFalse((self.root / 'Figures').exists())

    def test_cache_only_does_not_require_prepared_csvs(self):
        import shutil
        args = arguments(self.root); caches(args)
        shutil.rmtree(self.root / 'features')
        args.pipeline_mode, args.train_miafex, args.extract_miafex = 'full', 'yes', 'yes'
        with patch.object(figures, 'publication_figures', tiny_figures), patch.object(figures, 'statistical_figures', tiny_figures):
            manifest = core.run_report(args)
        self.assertEqual(manifest['optimization_calls'], 0)

    def test_seed_schedule_run_ids_and_ambiguous_science_are_rejected(self):
        import scientific_cache
        args = arguments(self.root); caches(args)
        path = next(self.root.rglob('*.pkl'))
        original = path.read_bytes()
        entry = pickle.loads(original)
        entry['row']['CompletedRunIDs'] = [0, 2, 3]
        path.write_bytes(pickle.dumps(entry))
        with self.assertRaisesRegex(ValueError, 'run IDs'): core.run_report(args)
        path.write_bytes(original)
        entry = pickle.loads(original)
        entry['identity']['optimizer_implementation']['revision'] = 'different-science'
        digest = scientific_cache.identity_digest(entry['identity'])
        path.with_name(f'{digest}_results.pkl').write_bytes(pickle.dumps(entry))
        with self.assertRaisesRegex(ValueError, 'Ambiguous scientific identity'): core.run_report(args)
        self.assertFalse((self.root / 'Figures').exists())
