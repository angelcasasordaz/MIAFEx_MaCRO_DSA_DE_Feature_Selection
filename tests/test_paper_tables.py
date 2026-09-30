"""Small synthetic manuscript workbooks; no experiments or figure rendering."""
from contextlib import ExitStack
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
import tempfile
import unittest
from unittest.mock import patch

import numpy as np
import pandas as pd
from openpyxl import load_workbook

import main_best as m
from reporting import paper_tables as paper


def fixture(datasets=("First", "Second"), algorithms=("DE", "PSO"), classifiers=("knn", "rf"), runs=3):
    args = SimpleNamespace(experiment_mode="full", estimators=list(classifiers),
                           optimizers=list(algorithms), transfer_functions=["vstf_01"],
                           full_replica_report_only=False, exp_id=913)
    results = {}
    for d, ds in enumerate(datasets):
        results[ds] = {}
        for a, opt in enumerate(algorithms):
            for c, cls in enumerate(classifiers):
                n = runs if isinstance(runs, int) else runs[a, c]
                base = np.arange(n, dtype=float) + d*2 + a + c
                results[ds][f"{opt}_{cls.upper()}"] = {
                    "Estimator": cls, "AccRuns": 60+base*2, "F1Runs": .5+base/100,
                    "PSRuns": .6+base/100, "RSRuns": .7+base/100,
                    "FitRuns": .3+base/100, "FeatRuns": 2+base, "TimeRuns": 1+base,
                }
    return args, results


class PaperTableTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.paths = SimpleNamespace(exp_tag="EXP913", res_dir=self.temp.name, fig_dir=self.temp.name)

    def export(self, args, results):
        path = paper.export_paper_tables(results, list(results), args.optimizers, args, self.paths)
        self.assertEqual(Path(path).name, "Paper_Tables_EXP913.xlsx")
        wb = load_workbook(path, data_only=True)
        self.addCleanup(wb.close)
        return wb

    def test_all_statistics_units_directions_and_per_run_averaging(self):
        args, results = fixture()
        wb = self.export(args, results)
        self.assertEqual(wb.sheetnames, ["Overall", "Datasets_1", "Datasets_2"])
        for a, opt in enumerate(args.optimizers):
            for metric_index, metric in enumerate(paper.METRICS):
                for c, cls in enumerate(args.estimators):
                    vectors = [np.asarray(records[f"{opt}_{cls.upper()}"][metric.run_key])/metric.scale
                               for records in results.values()]
                    overall = np.mean(vectors, axis=0)
                    expected = [max(overall), min(overall), np.mean(overall), np.std(overall, ddof=1)]
                    if metric.best_mode == "min": expected[:2] = expected[1::-1]
                    actual = [wb["Overall"].cell(3+4*a+s, 3+2*metric_index+c).value for s in range(4)]
                    np.testing.assert_allclose(actual, expected, rtol=1e-14)
                    for d, vector in enumerate(vectors):
                        expected = [max(vector), min(vector), np.mean(vector), np.std(vector, ddof=1)]
                        if metric.best_mode == "min": expected[:2] = expected[1::-1]
                        actual = [wb[f"Datasets_{c+1}"].cell(3+4*a+s, 3+7*d+metric_index).value for s in range(4)]
                        np.testing.assert_allclose(actual, expected, rtol=1e-14)
        self.assertEqual(wb["Overall"].freeze_panes, "C3")
        self.assertEqual(wb["Overall"]["C3"].number_format, "0.0000")

    def test_variable_dimensions_classifiers_metrics_and_run_counts(self):
        for datasets, algorithms, classifiers, runs in (
            (("Only",), ("DE",), ("tree",), 1),
            (("A", "B", "C", "D"), ("DE", "PSO", "JADE"), ("knn", "rf"),
             {(a, c): a+c+1 for a in range(3) for c in range(2)}),
        ):
            with self.subTest(datasets=datasets):
                args, results = fixture(datasets, algorithms, classifiers, runs)
                for records in results.values():
                    for row in records.values():
                        for metric in paper.METRICS[1:]: row.pop(metric.run_key)
                wb = self.export(args, results)
                self.assertEqual(wb["Overall"].max_row, len(algorithms)*4+2)
                self.assertEqual(wb["Overall"].max_column, len(classifiers)+2)
                self.assertEqual(len(wb.sheetnames), len(classifiers)+1)
                self.assertEqual(wb["Overall"]["C6"].value, 0)
                self.assertIn("run counts", wb.properties.description)

    def test_nonfinite_values_use_existing_reducer(self):
        args, results = fixture(datasets=("Only",), algorithms=("DE",), classifiers=("knn",))
        results["Only"]["DE_KNN"]["AccRuns"] = [np.nan, 80, np.inf]
        results["Only"]["DE_KNN"]["F1Runs"] = [np.nan]*3
        wb = self.export(args, results)
        self.assertEqual([wb["Overall"].cell(r, 3).value for r in range(3, 7)], [.8, .8, .8, 0])
        self.assertEqual([wb["Overall"].cell(r, 4).value for r in range(3, 7)], [None]*4)

    def test_incomplete_or_ambiguous_runs_are_rejected_without_writing(self):
        for problem in ("unequal", "missing", "duplicate", "empty"):
            args, results = fixture()
            row = results["First"]["DE_KNN"]
            if problem == "unequal": row["AccRuns"] = [80]
            if problem == "missing": del row["AccRuns"]
            if problem == "duplicate": results["First"]["DE_VSTF_01_KNN"] = row
            if problem == "empty": row["AccRuns"] = []
            with self.subTest(problem=problem), self.assertRaises(ValueError):
                self.export(args, results)
        self.assertFalse(list(Path(self.temp.name).glob("*.xlsx")))

    def test_report_only_uses_indexed_exporter(self):
        args, results = fixture()
        args.report_only = True
        with self.assertRaisesRegex(ValueError, 'export_indexed_tables'):
            self.export(args, results)


if __name__ == "__main__":
    unittest.main()
